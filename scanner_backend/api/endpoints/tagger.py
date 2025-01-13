import codecs
import logging
import pickle
import subprocess
from pdfminer.high_level import extract_text as extract_pdf_text
from docx import Document
from pptx import Presentation
from os.path import splitext
import tempfile

from flask import request, abort
from flask_restx import Namespace, Resource

import tipi_tasks
from scanner_backend.api.business import get_tags
from scanner_backend.api.lines_of_action import LinesOfAction
from scanner_backend.api.endpoints import cache, limiter
from scanner_backend.api.parsers import parser_tagger
from scanner_backend.settings import Config


log = logging.getLogger(__name__)

ns = Namespace(
    "tagger", description="Operations related to tag texts using our knowledge base"
)


@ns.route("/")
@ns.expect(parser_tagger)
class TaggerExtractor(Resource):

    def post(self):
        """Returns a list of topics and tags matching the text."""
        try:
            cache_key = Config.CACHE_TAGS
            tags = cache.get(cache_key)
            if tags is None:
                tags = get_tags()
                cache.set(cache_key, tags, timeout=5 * 60)
            tags = codecs.encode(pickle.dumps(tags), "base64").decode()
            tipi_tasks.init()
            text = ""
            if "text" in request.form and request.form["text"]:
                text = request.form["text"]
            else:
                if "file" in request.files:
                    file_input = request.files["file"]
                    with tempfile.NamedTemporaryFile(
                        prefix="tipiscanner_", suffix=splitext(file_input.filename)[1]
                    ) as f:
                        f.write(file_input.stream.read())
                        f.seek(0)
                        print("MIMETYPE:", file_input.mimetype)
                        if file_input.mimetype == "text/plain":
                            text = f.read().decode("utf-8").strip()
                        elif file_input.mimetype == "application/pdf":
                            text = extract_pdf_text(f.name).strip()
                        elif (
                            file_input.mimetype
                            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        ):
                            doc = Document(f)
                            text = "\n".join(
                                [para.text for para in doc.paragraphs]
                            ).strip()
                        elif file_input.mimetype == "application/msword":
                            result = subprocess.run(
                                ["antiword", f.name],
                                stdout=subprocess.PIPE,
                                stderr=subprocess.PIPE,
                            )
                            if result.returncode != 0:
                                raise Exception(
                                    f"Error al leer el archivo .doc: {result.stderr.decode('utf-8')}"
                                )
                            text = result.stdout.decode("utf-8").strip()
                        elif (
                            file_input.mimetype
                            == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                        ):
                            ppt = Presentation(f)
                            text = "\n".join(
                                [
                                    shape.text
                                    for slide in ppt.slides
                                    for shape in slide.shapes
                                    if hasattr(shape, "text")
                                ]
                            ).strip()
                        else:
                            abort(
                                400,
                                "Formato no soportado. Por favor, utilice un archivo .txt, .pdf, .docx, .doc o .pptx.",
                            )
                        f.close()
                    if not text:
                        abort(
                            400,
                            "Error al obtener el texto del fichero proporcionado. Pruebe con otro fichero.",
                        )
            text_length = len(text.split())

            if text_length >= Config.TAGGER_MAX_WORDS:
                task = tipi_tasks.tagger.extract_tags_from_text.apply_async(
                    (text, tags)
                )
                eta_time = int((text_length / 1000) * 4)
                task_id = task.id
                result = {
                    "status": "PROCESSING",
                    "task_id": task_id,
                    "estimated_time": eta_time,
                }
            else:
                result = tipi_tasks.tagger.extract_tags_from_text(text, tags)
                LinesOfAction.extract(result)
            return result
        except Exception as e:
            if hasattr(e, "code") and hasattr(e, "description"):
                abort(e.code, e.description)
            else:
                abort(500, "Internal server error")


@ns.route("/result/<id>")
@ns.param(
    name="id",
    description="Task id",
    type=str,
    required=True,
    location=["path"],
    help="Invalid identifier",
)
@ns.response(404, "Task not found.")
class TaggerResult(Resource):

    def get(self, id):
        """Returns tagging task's result"""
        tipi_tasks.init()
        return tipi_tasks.tagger.check_status_task(id)
